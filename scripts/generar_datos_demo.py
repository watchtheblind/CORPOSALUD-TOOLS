"""Genera Excel de ejemplo para probar el combinador.

Los archivos se crean en data/entrada/ y NO se versionan (data/ está en .gitignore).
"""

from pathlib import Path

from openpyxl import Workbook, load_workbook

RAIZ = Path(__file__).resolve().parent.parent
DIR_ENTRADA = RAIZ / "data" / "entrada"


def escribir(nombre: str, filas: list[list[str | None]]) -> None:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "DATOS"
    for fila in filas:
        hoja.append(fila)
    libro.save(DIR_ENTRADA / nombre)
    print(f"creado: {DIR_ENTRADA / nombre}")


def agregar_hoja(nombre: str, hoja: str, filas: list[list[str | None]]) -> None:
    ruta = DIR_ENTRADA / nombre
    libro = load_workbook(ruta)
    nueva = libro.create_sheet(hoja)
    for fila in filas:
        nueva.append(fila)
    libro.save(ruta)
    print(f"hoja {hoja!r} agregada a: {ruta}")


def main() -> None:
    DIR_ENTRADA.mkdir(parents=True, exist_ok=True)

    # Archivo de referencia: define el encabezado "oficial".
    escribir(
        "demo_pacientes_enero.xlsx",
        [
            ["CEDULA", "NOMBRE", "FECHA", "SEDE", "EDAD"],
            ["01234567", "Ana Pérez", "2026-01-05", "Caracas", "34"],
            ["12345678", "Luis Gómez", "2026-01-07", "Valencia", "41"],
            ["23456789", "María Rodríguez", "2026-01-12", "Maracaibo", "29"],
        ],
    )

    # Mismos datos, títulos distintos + una columna extra (CORREO).
    escribir(
        "demo_pacientes_febrero.xlsx",
        [
            ["C.I.", "NOMBRE COMPLETO", "FECHA_ATENCION", "SEDE", "EDAD", "CORREO"],
            ["01234567", "Ana Pérez", "2026-02-03", "Caracas", "34", "ana@example.com"],
            ["34567890", "Pedro Díaz", "2026-02-08", "Barquisimeto", "52", "pedro@example.com"],
        ],
    )

    # Columnas en otro orden + otra columna extra (OBSERVACION) + fila vacía al final.
    escribir(
        "demo_pacientes_marzo.xlsx",
        [
            ["SEDE", "CEDULA", "EDAD", "NOMBRES Y APELLIDOS", "FECHA", "OBSERVACION"],
            ["Maracay", "45678901", "38", "José Martínez", "2026-03-02", "control"],
            ["Caracas", "01234567", "34", "Ana Pérez", "2026-03-09", "seguimiento"],
            [None, None, None, None, None, None],
        ],
    )

    # Segunda hoja en marzo, para probar la selección de hoja por archivo.
    agregar_hoja(
        "demo_pacientes_marzo.xlsx",
        "HOJA_ALTERNATIVA",
        [
            ["SEDE", "CEDULA", "EDAD", "NOMBRES Y APELLIDOS", "FECHA", "OBSERVACION"],
            ["Maracay", "45678901", "38", "José Martínez", "2026-03-02", "control"],
            ["Caracas", "56789012", "45", "Lucía Fernández", "2026-03-15", "primera vez"],
        ],
    )


if __name__ == "__main__":
    main()
