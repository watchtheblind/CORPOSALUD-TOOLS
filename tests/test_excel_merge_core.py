"""Pruebas de la lógica de consolidación, sin interfaz gráfica."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from excel_merge_core import combinar_excels, escanear_hojas


def crear_excel(ruta: Path, filas: list[list[object]]) -> Path:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "DATOS"
    for fila in filas:
        hoja.append(fila)
    libro.save(ruta)
    return ruta


class CombinarExcelsTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

        self.ref = crear_excel(
            self.dir / "ref.xlsx",
            [
                ["CEDULA", "NOMBRE", "FECHA", "SEDE"],
                ["01234567", "Ana", "2026-01-01", "Caracas"],
                ["12345678", "Luis", "2026-01-02", "Valencia"],
            ],
        )
        self.feb = crear_excel(
            self.dir / "feb.xlsx",
            [
                ["C.I.", "NOMBRE COMPLETO", "FECHA_ATENCION", "SEDE", "CORREO"],
                ["01234567", "Ana", "2026-02-01", "Caracas", "ana@example.com"],
            ],
        )
        self.mar = crear_excel(
            self.dir / "mar.xlsx",
            [
                ["SEDE", "CEDULA", "NOMBRES Y APELLIDOS", "FECHA", "OBSERVACION"],
                ["Maracay", "45678901", "José", "2026-03-01", "control"],
            ],
        )
        self.ajustes = {
            "feb.xlsx": {
                "CEDULA": "C.I.",
                "NOMBRE": "NOMBRE COMPLETO",
                "FECHA": "FECHA_ATENCION",
            },
            "mar.xlsx": {
                "NOMBRE": "NOMBRES Y APELLIDOS",
            },
        }

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def combinar(self, **kwargs):
        parametros = {
            "rutas": [self.ref, self.feb, self.mar],
            "archivo_referencia": "ref.xlsx",
            "ajustes": self.ajustes,
        }
        parametros.update(kwargs)
        return combinar_excels(**parametros)

    def test_usa_encabezado_de_la_referencia(self) -> None:
        df = self.combinar()
        self.assertEqual(
            list(df.columns),
            [
                "CEDULA",
                "NOMBRE",
                "FECHA",
                "SEDE",
                "archivo_origen",
                "hoja_origen",
                "fila_origen",
            ],
        )
        self.assertEqual(len(df), 4)

    def test_mapea_columnas_con_nombres_distintos(self) -> None:
        df = self.combinar(extras={"feb.xlsx": ["CORREO"]})
        fila_feb = df[df["archivo_origen"] == "feb.xlsx"].iloc[0]
        self.assertEqual(fila_feb["CEDULA"], "01234567")
        self.assertEqual(fila_feb["NOMBRE"], "Ana")
        self.assertEqual(fila_feb["FECHA"], "2026-02-01")
        self.assertEqual(fila_feb["CORREO"], "ana@example.com")

    def test_descarta_encabezado_y_columnas_no_aprobadas(self) -> None:
        df = self.combinar(extras={"feb.xlsx": ["CORREO"]})
        self.assertNotIn("OBSERVACION", df.columns)
        self.assertNotIn("C.I.", df.columns)
        self.assertNotIn("NOMBRES Y APELLIDOS", df.columns)

    def test_preserva_ceros_a_la_izquierda(self) -> None:
        df = self.combinar()
        self.assertIn("01234567", df["CEDULA"].tolist())
        self.assertNotIn("1234567", df["CEDULA"].tolist())

    def test_trazabilidad(self) -> None:
        df = self.combinar()
        fila_mar = df[df["archivo_origen"] == "mar.xlsx"].iloc[0]
        self.assertEqual(fila_mar["hoja_origen"], "DATOS")
        self.assertEqual(fila_mar["fila_origen"], 2)  # primera fila de datos
        self.assertEqual(fila_mar["SEDE"], "Maracay")

    def test_columna_sin_mapear_queda_vacia(self) -> None:
        df = self.combinar(ajustes={}, usar_sugerencias=False)
        self.assertTrue((df["CEDULA"] == "").all())
        self.assertTrue((df["NOMBRE"] == "").all())

    def test_error_si_se_exige_mapeo_completo(self) -> None:
        with self.assertRaises(ValueError):
            self.combinar(
                ajustes={},
                usar_sugerencias=False,
                exigir_mapeo_completo=True,
            )

    def test_error_si_una_columna_fuente_se_usa_dos_veces(self) -> None:
        ajustes = {
            "feb.xlsx": {
                "CEDULA": "C.I.",
                "NOMBRE": "C.I.",
            }
        }
        with self.assertRaises(ValueError):
            self.combinar(ajustes=ajustes)

    def test_error_si_el_encabezado_tiene_duplicados(self) -> None:
        malo = crear_excel(
            self.dir / "malo.xlsx",
            [["CEDULA", "CEDULA"], ["1", "2"]],
        )
        with self.assertRaises(ValueError):
            combinar_excels([malo], "malo.xlsx")

    def test_guarda_el_archivo_de_salida(self) -> None:
        salida = self.dir / "salida" / "consolidado.xlsx"
        df = self.combinar(guardar_en=salida)
        self.assertTrue(salida.exists())

        import pandas as pd

        releido = pd.read_excel(salida, dtype=str, keep_default_na=False)
        self.assertEqual(releido.shape, df.shape)

    def test_escanear_hojas(self) -> None:
        multi = self.dir / "multi.xlsx"
        libro = Workbook()
        libro.active.title = "RESUMEN"
        libro.create_sheet("DATOS")
        libro.save(multi)

        self.assertEqual(
            escanear_hojas([multi])["multi.xlsx"],
            ["RESUMEN", "DATOS"],
        )

    def test_elige_hoja_de_origen_por_archivo(self) -> None:
        multi = self.dir / "multi.xlsx"
        libro = Workbook()
        resumen = libro.active
        resumen.title = "RESUMEN"
        resumen.append(["TEXTO"])
        resumen.append(["no usar"])
        datos = libro.create_sheet("DATOS")
        datos.append(["CEDULA", "NOMBRE"])
        datos.append(["999", "Desde hoja 2"])
        libro.save(multi)

        df = combinar_excels(
            [multi],
            "multi.xlsx",
            hojas_por_archivo={"multi.xlsx": "DATOS"},
        )
        self.assertEqual(df["NOMBRE"].tolist(), ["Desde hoja 2"])
        self.assertEqual(df["hoja_origen"].tolist(), ["DATOS"])

    def test_error_si_la_hoja_elegida_no_existe(self) -> None:
        multi = self.dir / "multi.xlsx"
        libro = Workbook()
        libro.active.title = "RESUMEN"
        libro.save(multi)

        with self.assertRaises(ValueError):
            combinar_excels(
                [multi],
                "multi.xlsx",
                hojas_por_archivo={"multi.xlsx": "NO_EXISTE"},
            )


if __name__ == "__main__":
    unittest.main()
